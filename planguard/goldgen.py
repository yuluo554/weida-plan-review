# -*- coding: utf-8 -*-
"""配对真值数据生成器核心（M2，设计见 plan/04 §2）。

程序化生成深基坑专项施工方案 docx（合成样例，不含真实项目信息）+ truth.json
（参数真值 + 注入差异清单），用于解析 F1 与端到端校核基准——评测零 API 依赖可重复。

注入差异类型（每份 1–2 类，count 份内轮转保证每类充分出现）。
基准不变量：自然参数不触发任何非 pass 结论；每处注入在 expect 记录主期望、
在 also_expect 记录同一注入隐含的其余非 pass 结论（真值=该文档全部非 pass 集合）：
- over_disp       位移控制值取 35mm（>30mm 示例上限）      → 期望 R-DP-021 fail
- low_freq        监测频率 0.5 次/d（<1 下限、<2 建议值）    → 期望 R-DP-031 fail、R-DP-061 fail
- multi_disp      计算书与监测方案给出两个不同位移控制值     → 多值冲突：该参数全部规则
                    （R-DP-021/022/062/073）均输出 manual
- unit_error      开挖深度放大 100 倍（超物理范围，应丢弃）   → 期望 R-DP-001 fail
- missing_section 缺"应急处置措施"章节                      → 期望 R-DP-051（章节齐套）fail
"""
import json
import random
from pathlib import Path
from typing import Any, Dict, List, Tuple

INJECTION_TYPES = ("over_disp", "low_freq", "multi_disp", "unit_error", "missing_section")
FS_INJECTION_TYPES = ("over_height", "low_spacing", "multi_height", "unit_error", "missing_section")

DOC_TITLE = "深基坑土方开挖及支护专项施工方案（合成样例）"
SUPPORT_TYPES = ("放坡+土钉墙", "桩锚支护", "地下连续墙+内支撑", "钢板桩+内支撑")
FS_DOC_TITLE = "高大模板支撑体系（高支模）专项施工方案（合成样例）"
FORMWORK_TYPES = ("扣件式钢管", "盘扣式钢管", "碗扣式钢管", "门式钢管")


def make_spec(rng: random.Random, idx: int, injections: List[str]) -> Dict[str, Any]:
    """随机抽参数 + 应用注入差异，返回文档内容与真值 spec。"""
    depth = round(rng.uniform(4.0, 9.0), 1)               # m
    disp = int(rng.uniform(20, 29))                        # mm（自然范围内不触发示例阈值30）
    freq = rng.choice([2.0, 2.0, 2.5])                     # 次/d（≥2，自然通过 R-DP-031/R-DP-061）
    support = rng.choice(SUPPORT_TYPES)
    sections: List[Tuple[str, List[str]]] = [
        ("一、工程概况", [
            "本工程为合成样例项目 G%04d（虚构，仅用于测试），基坑采用%s支护形式。"
            % (idx, support),
            "基坑开挖深度为%gm，基坑侧壁安全等级为二级。" % depth,
        ]),
        ("二、编制依据", [
            "1.《危险性较大的分部分项工程安全管理规定》（住建部令第37号）；",
            "2.《关于实施〈危险性较大的分部分项工程安全管理规定〉有关问题的通知》（建办质〔2018〕31号）；",
            "3.《建筑基坑支护技术规程》JGJ 120-2012；",
            "4.《建筑基坑工程监测技术标准》GB 50497-2019。",
        ]),
        ("三、施工计划", [
            "计划工期%d天，分%d段组织开挖。" % (rng.randint(30, 90), rng.randint(2, 4)),
        ]),
        ("四、施工工艺技术", [
            "4.1 土方开挖",
            "土方开挖分层分段进行，每层开挖深度不超过%gm。" % rng.choice([1.0, 1.5, 2.0]),
            "4.2 支护施工",
            "支护采用%s工艺，随挖随支，严禁超挖。" % support,
        ]),
        ("五、施工安全保证措施", [
            "5.1 组织保障",
            "成立以项目经理为组长的安全领导小组，责任到人。",
            "5.2 监测方案",
            "支护结构顶部水平位移控制值为%dmm。" % disp,
            "开挖至坑底期间，监测频率为%g次/d。" % freq,
        ]),
        ("六、施工管理及作业人员配备", [
            "现场配备专职安全员2名，特种作业人员持证上岗。",
        ]),
        ("七、验收要求", [
            "支护结构验收执行JGJ 120-2012相关规定，验收合格后方可进入下道工序。",
        ] + (["本工程基坑属超过一定规模的危大工程，已按规定组织专家论证，论证意见已落实。"]
             if depth >= 5 else [])),
        ("八、应急处置措施", [
            "编制应急预案，配备应急物资，定期组织演练。",
        ]),
        ("九、计算书及相关图纸", [
            "支护结构计算书及基坑支护平面图见附件（合成样例略）。",
        ]),
    ]

    expected: Dict[str, List[Any]] = {
        "dp.excavation_depth": [depth],
        "dp.retaining.displacement": [float(disp)],
        "dp.monitoring.frequency": [float(freq)],
        "dp.support_type": [support],
        "dp.safety_grade": ["二级"],
    }
    injected: List[Dict[str, Any]] = []

    def _paras(prefix: str) -> List[str]:
        """按标题前缀定位章节的段落列表（章节可能已被 missing_section 删除）。"""
        for heading, paras in sections:
            if heading.startswith(prefix):
                return paras
        raise KeyError("章节不存在: " + prefix)

    for kind in injections:
        if kind == "over_disp":
            disp = 35
            _paras("五、")[3] = "支护结构顶部水平位移控制值为35mm。"
            expected["dp.retaining.displacement"] = [35.0]
            injected.append({"type": kind,
                             "detail": "位移控制值取35mm，超过示例上限30mm",
                             "expect": {"rule_id": "R-DP-021", "result": "fail"}})
        elif kind == "low_freq":
            freq = 0.5
            _paras("五、")[4] = "开挖至坑底期间，监测频率为0.5次/d。"
            expected["dp.monitoring.frequency"] = [0.5]
            injected.append({"type": kind,
                             "detail": "监测频率0.5次/d，低于下限1次/d与建议值2次/d",
                             "expect": {"rule_id": "R-DP-031", "result": "fail"},
                             "also_expect": [{"rule_id": "R-DP-061", "result": "fail"}]})
        elif kind == "multi_disp":
            disp2 = disp + 5
            _paras("九、").append("计算书中支护结构顶部水平位移控制值取%dmm。" % disp2)
            expected["dp.retaining.displacement"] = sorted({float(disp), float(disp2)})
            injected.append({"type": kind,
                             "detail": "监测方案%imm 与计算书%imm 两处位移控制值冲突" % (disp, disp2),
                             "expect": {"rule_id": "R-DP-021", "result": "manual"},
                             "also_expect": [
                                 {"rule_id": "R-DP-022", "result": "manual"},
                                 {"rule_id": "R-DP-062", "result": "manual"},
                                 {"rule_id": "R-DP-073", "result": "manual"},
                             ]})
        elif kind == "unit_error":
            _paras("一、")[1] = "基坑开挖深度为%gm，基坑侧壁安全等级为二级。" % (depth * 100)
            expected.pop("dp.excavation_depth", None)  # 超物理范围应被丢弃 → 无该参数卡
            injected.append({"type": kind,
                             "detail": "开挖深度误写为%gm（超物理范围，提取器应丢弃）" % (depth * 100),
                             "expect": {"rule_id": "R-DP-001", "result": "fail"}})
        elif kind == "missing_section":
            sections = [s for s in sections if s[0] != "八、应急处置措施"]
            injected.append({"type": kind,
                             "detail": "缺少「应急处置措施」章节",
                             "expect": {"rule_id": "R-DP-051", "result": "fail", "since": "M3"}})
    return {
        "idx": idx, "title": DOC_TITLE, "sections": sections,
        "expected_cards": expected, "injected": injected,
    }


def make_fs_spec(rng: random.Random, idx: int, injections: List[str]) -> Dict[str, Any]:
    """高支模（模板支撑体系）spec：参数与注入逻辑同 make_spec，口径不变。

    自然参数：搭设高度 5.5–7.5m（≥5m 属危大须专项方案，<8m 不触发论证门槛）、
    立杆间距 0.8–1.1m（≤示例上限 1.2m）；论证句措辞避开"别名+数字"（HANDOFF-M4 问题 A）。
    """
    height = round(rng.uniform(5.5, 7.5), 1)               # m
    spacing = rng.choice([0.8, 0.9, 0.9, 1.0, 1.1])        # m
    formwork = rng.choice(FORMWORK_TYPES)
    sections: List[Tuple[str, List[str]]] = [
        ("一、工程概况", [
            "本工程为合成样例项目 G%04d（虚构，仅用于测试），模板支撑体系采用%s脚手架搭设。"
            % (idx, formwork),
            "本工程属危大工程中的混凝土模板支撑工程，编制本专项施工方案。",
        ]),
        ("二、编制依据", [
            "1.《危险性较大的分部分项工程安全管理规定》（住建部令第37号）；",
            "2.《关于实施〈危险性较大的分部分项工程安全管理规定〉有关问题的通知》（建办质〔2018〕31号）；",
            "3.《建筑施工模板安全技术规范》JGJ 162-2008；",
            "4.《建筑施工扣件式钢管脚手架安全技术规范》JGJ 130-2011。",
        ]),
        ("三、施工计划", [
            "计划工期%d天，分%d区组织搭设与浇筑。" % (rng.randint(20, 60), rng.randint(2, 4)),
        ]),
        ("四、施工工艺技术", [
            "4.1 模板支撑体系设计",
            "模板支撑体系搭设高度为%gm。" % height,
            "支撑体系采用%s钢管，立杆底部设置垫板与扫地杆。" % formwork,
            "4.2 混凝土浇筑",
            "混凝土分层浇筑，浇筑顺序按方案执行，严禁超载堆放。",
        ]),
        ("五、施工安全保证措施", [
            "5.1 组织保障",
            "成立以项目经理为组长的安全领导小组，责任到人。",
            "5.2 监测方案",
            "模板支撑体系立杆间距为%gm，浇筑期间安排专人对支撑体系变形进行监测。" % spacing,
        ]),
        ("六、施工管理及作业人员配备", [
            "现场配备专职安全员2名，特种作业人员持证上岗。",
        ]),
        ("七、验收要求", [
            "支撑体系验收执行JGJ 162-2008相关规定，验收合格后方可浇筑混凝土。",
        ]),
        ("八、应急处置措施", [
            "编制应急预案，配备应急物资，定期组织演练。",
        ]),
        ("九、计算书及相关图纸", [
            "模板支撑体系计算书及模板平面图见附件（合成样例略）。",
        ]),
    ]

    expected: Dict[str, List[Any]] = {
        "fs.support.height": [height],
        "fs.pole.spacing": [float(spacing)],
        # 提取器捕获"钢管"前的短名（如"扣件式"），真值与 R-FS-041 enum 同口径
        "fs.formwork.type": [formwork.replace("钢管", "")],
    }
    injected: List[Dict[str, Any]] = []

    def _paras(prefix: str) -> List[str]:
        for heading, paras in sections:
            if heading.startswith(prefix):
                return paras
        raise KeyError("章节不存在: " + prefix)

    for kind in injections:
        if kind == "over_height":
            # 若与 multi_height 组合，先移除计算书高度行，避免两注入互相破坏期望
            for heading, paras in sections:
                if heading.startswith("九、"):
                    paras[:] = [p for p in paras
                                if not p.startswith("计算书中模板支撑体系搭设高度取")]
            height = 8.5
            _paras("四、")[1] = "模板支撑体系搭设高度为8.5m。"
            expected["fs.support.height"] = [8.5]
            injected.append({"type": kind,
                             "detail": "搭设高度取8.5m，达到超过一定规模门槛8m且全文未体现专家论证",
                             "expect": {"rule_id": "R-FS-101", "result": "fail"}})
        elif kind == "low_spacing":
            spacing = 1.5
            _paras("五、")[3] = "模板支撑体系立杆间距为1.5m，浇筑期间安排专人对支撑体系变形进行监测。"
            expected["fs.pole.spacing"] = [1.5]
            injected.append({"type": kind,
                             "detail": "立杆间距1.5m，超过示例上限1.2m",
                             "expect": {"rule_id": "R-FS-031", "result": "fail"}})
        elif kind == "multi_height":
            height2 = round(height + 1.0, 1)
            if height2 >= 8.0:  # 第二值不得达到 8m 论证门槛，否则误触发 R-FS-101
                height2 = 7.8
            _paras("九、").append("计算书中模板支撑体系搭设高度取%gm。" % height2)
            expected["fs.support.height"] = sorted({float(height), height2})
            injected.append({"type": kind,
                             "detail": "方案%gm 与计算书%gm 两处搭设高度冲突" % (height, height2),
                             "expect": {"rule_id": "R-FS-012", "result": "manual"}})
        elif kind == "unit_error":
            _paras("五、")[3] = "模板支撑体系立杆间距为%gm，浇筑期间安排专人对支撑体系变形进行监测。" % (spacing * 100)
            expected.pop("fs.pole.spacing", None)
            injected.append({"type": kind,
                             "detail": "立杆间距误写为%gm（超物理范围，提取器应丢弃）" % (spacing * 100),
                             "expect": {"rule_id": "R-FS-003", "result": "fail"}})
        elif kind == "missing_section":
            sections = [s for s in sections if s[0] != "八、应急处置措施"]
            injected.append({"type": kind,
                             "detail": "缺少「应急处置措施」章节",
                             "expect": {"rule_id": "R-FS-051", "result": "fail"}})
    return {
        "idx": idx, "title": FS_DOC_TITLE, "sections": sections,
        "expected_cards": expected, "injected": injected,
    }


def write_sample(spec: Dict[str, Any], docx_path: Path) -> None:
    """按 spec 生成 docx（需 python-docx）。标题用 Heading 样式便于章节识别。"""
    import docx  # python-docx

    document = docx.Document()
    document.add_heading(spec["title"], level=0)
    for heading, paragraphs in spec["sections"]:
        document.add_heading(heading, level=1)
        for text in paragraphs:
            document.add_paragraph(text)
    document.save(str(docx_path))


def generate(out_dir: Path, gold_dir: Path, count: int = 10, seed: int = 2026,
             category: str = "deep_pit") -> List[Dict[str, Any]]:
    """生成 count 份样例 + 真值；注入类型按轮转分布，保证每类充分出现。

    category: deep_pit（文件名 gen_XXXX）| formwork_support（文件名 fs_XXXX）。
    """
    try:
        import docx  # noqa: F401  提前给出可读错误
    except ImportError as exc:
        raise RuntimeError(
            "生成 docx 需要 python-docx。请先安装：\n"
            "  py -m pip install -i https://pypi.tuna.tsinghua.edu.cn/simple python-docx"
        ) from exc

    if category not in ("deep_pit", "formwork_support"):
        raise ValueError("未知类目: %s（支持 deep_pit / formwork_support）" % category)
    make_spec_fn = make_spec if category == "deep_pit" else make_fs_spec
    injection_types = INJECTION_TYPES if category == "deep_pit" else FS_INJECTION_TYPES

    out_dir, gold_dir = Path(out_dir), Path(gold_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    gold_dir.mkdir(parents=True, exist_ok=True)

    rng = random.Random(seed)
    prefix = "gen" if category == "deep_pit" else "fs"
    truths: List[Dict[str, Any]] = []
    for i in range(count):
        injections: List[str] = []
        for slot in range(2):  # 每份注入 1–2 类，去重
            kind = injection_types[(i + slot * (count // 2 + 1)) % len(injection_types)]
            if kind not in injections:
                injections.append(kind)
        spec = make_spec_fn(rng, i + 1, injections)
        name = "%s_%04d" % (prefix, i + 1)
        docx_path = out_dir / (name + ".docx")
        write_sample(spec, docx_path)
        truth = {
            "file": docx_path.name,
            "category": category,
            "seed": seed,
            "title": spec["title"],
            "expected_cards": spec["expected_cards"],
            "injected": spec["injected"],
        }
        (gold_dir / (name + ".truth.json")).write_text(
            json.dumps(truth, ensure_ascii=False, indent=2), encoding="utf-8")
        truths.append(truth)
    return truths
