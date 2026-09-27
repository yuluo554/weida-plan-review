# -*- coding: utf-8 -*-
"""报告导出（M3 Markdown / M5 docx）。

报告结构（plan/03 §7）：总体结论 → 不合规清单 → 待人工确认 →
通过项 → 参数卡附录 → 签署栏。报告不替代专家论证。
"""
from datetime import datetime
from pathlib import Path

from .. import __version__

RESULT_LABEL = {"fail": "不合规", "manual": "待人工确认", "pass": "通过"}
RESULT_ORDER = {"fail": 0, "manual": 1, "pass": 2}


def _overall(summary) -> str:
    if summary["fail"]:
        return "**总体结论: 不合规** —— 存在强制条款不满足，须整改后复审。"
    if summary["manual"]:
        return ("**总体结论: 待人工确认** —— 自动核查未发现违规，"
                "但有 %d 项需人工复核。" % summary["manual"])
    return "**总体结论: 通过**（基于当前规则库）。"


def _finding_table(findings) -> list:
    rows = [
        "| 规则 | 级别 | 结论 | 说明 | 依据条款 | 整改建议 | 原文证据 |",
        "| --- | --- | --- | --- | --- | --- | --- |",
    ]
    for f in sorted(findings, key=lambda x: (RESULT_ORDER.get(x.result, 9), x.rule_id)):
        rows.append("| %s %s | %s | %s | %s | %s | %s | %s |" % (
            f.rule_id, f.rule_name, f.level, RESULT_LABEL.get(f.result, f.result),
            f.detail or "-", f.basis or "-", f.advice or "-",
            (f.evidence_line or "-").replace("|", "\\|"),
        ))
    return rows


def export_markdown(result, out_path) -> Path:
    """把 ReviewResult 导出为 Markdown 审查报告，返回写入路径。"""
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    s = result.summary()

    lines = [
        "# PlanGuard 专项施工方案审查报告",
        "",
        "- 审查对象: %s" % result.doc,
        "- 审查时间: %s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "- 工具版本: planguard %s" % __version__,
        "- 规则库规模: %s 条" % result.meta.get("rules_total", "?"),
        "",
        "## 一、总体结论",
        "",
        "| 结论 | 数量 |",
        "| --- | --- |",
        "| 不合规 | %d |" % s["fail"],
        "| 待人工确认 | %d |" % s["manual"],
        "| 通过 | %d |" % s["pass"],
        "",
        _overall(s),
        "",
        "## 二、不合规清单",
        "",
    ]
    fails = [f for f in result.findings if f.result == "fail"]
    lines += _finding_table(fails) if fails else ["（无）"]
    lines += ["", "## 三、待人工确认清单", ""]
    manuals = [f for f in result.findings if f.result == "manual"]
    lines += _finding_table(manuals) if manuals else ["（无）"]
    lines += ["", "## 四、通过项", ""]
    passes = [f for f in result.findings if f.result == "pass"]
    lines += _finding_table(passes) if passes else ["（无）"]
    lines += ["", "## 五、附录：参数卡全表", "",
              "| 参数 | 取值 | 单位/类型 | 出处章节 | 原文 |", "| --- | --- | --- | --- | --- |"]
    for c in result.cards:
        val = c.text_value if c.text_value is not None else c.value
        unit = c.unit or ("文本" if c.text_value is not None else "")
        ev = c.evidence
        lines.append("| %s %s | %s | %s | %s | %s |" % (
            c.param_id, c.name, val, unit,
            ev.section if ev else "-",
            (ev.line_text if ev else "-").replace("|", "\\|")))
    lines += [
        "",
        "## 六、签署栏",
        "",
        "| 审查人 | 日期 | 备注 |",
        "| --- | --- | --- |",
        "|  |  | 本报告为辅助预审，不替代专家论证 |",
        "",
    ]
    discarded = result.meta.get("discarded") or []
    if discarded:
        lines += ["## 附：已丢弃的可疑取值（超出物理合理范围）", ""]
        for d in discarded:
            lines.append("- %s = %s%s（%s）原文：%s" % (
                d.get("param"), d.get("value"), d.get("unit", ""), d.get("reason"),
                d.get("line", "")))
        lines.append("")
    out_path.write_text("\n".join(lines), encoding="utf-8")
    return out_path


def export_docx(result, out_path) -> Path:
    """把 ReviewResult 导出为可归档的 docx 审查报告（M5，python-docx）。

    结构同 Markdown 版（plan/03 §7）：封面信息 → 总体结论 → 不合规清单 →
    待人工确认 → 通过项 → 参数卡附录 → 签署栏（+丢弃值附表）。
    """
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise RuntimeError(
            "导出 docx 报告需要 python-docx。请先安装：\n"
            "  py -m pip install -i https://pypi.tuna.tsinghua.edu.cn/simple python-docx"
        ) from exc

    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    s = result.summary()

    document.add_heading("PlanGuard 专项施工方案审查报告", level=0)
    for text in (
        "审查对象: %s" % result.doc,
        "审查时间: %s" % datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "工具版本: planguard %s" % __version__,
        "规则库规模: %s 条" % result.meta.get("rules_total", "?"),
    ):
        document.add_paragraph(text)

    document.add_heading("一、总体结论", level=1)
    table = document.add_table(rows=1, cols=2)
    table.style = "Light Grid Accent 1"
    _fill_row(table.rows[0], ("结论", "数量"))
    for label, key in (("不合规", "fail"), ("待人工确认", "manual"), ("通过", "pass")):
        _fill_row(table.add_row(), (label, str(s[key])))
    document.add_paragraph(_overall_plain(s))

    fails = [f for f in result.findings if f.result == "fail"]
    manuals = [f for f in result.findings if f.result == "manual"]
    passes = [f for f in result.findings if f.result == "pass"]
    document.add_heading("二、不合规清单", level=1)
    _findings_docx(document, fails)
    document.add_heading("三、待人工确认清单", level=1)
    _findings_docx(document, manuals)
    document.add_heading("四、通过项", level=1)
    _findings_docx(document, passes)

    document.add_heading("五、附录：参数卡全表", level=1)
    table = document.add_table(rows=1, cols=5)
    table.style = "Light Grid Accent 1"
    _fill_row(table.rows[0], ("参数", "取值", "单位/类型", "出处章节", "原文"))
    for c in result.cards:
        val = c.text_value if c.text_value is not None else c.value
        unit = c.unit or ("文本" if c.text_value is not None else "")
        ev = c.evidence
        _fill_row(table.add_row(), (
            "%s %s" % (c.param_id, c.name), "%s" % val, unit,
            ev.section if ev else "-", ev.line_text if ev else "-"))

    document.add_heading("六、签署栏", level=1)
    table = document.add_table(rows=2, cols=3)
    table.style = "Light Grid Accent 1"
    _fill_row(table.rows[0], ("审查人", "日期", "备注"))
    _fill_row(table.rows[1], ("", "", "本报告为辅助预审，不替代专家论证"))

    discarded = result.meta.get("discarded") or []
    if discarded:
        document.add_heading("附：已丢弃的可疑取值（超出物理合理范围）", level=1)
        for d in discarded:
            document.add_paragraph("%s = %s%s（%s）原文：%s" % (
                d.get("param"), d.get("value"), d.get("unit", ""), d.get("reason"),
                d.get("line", "")))
    document.save(str(out_path))
    return out_path


def _fill_row(row, cells) -> None:
    for cell, text in zip(row.cells, cells):
        cell.text = str(text)


def _overall_plain(summary) -> str:
    if summary["fail"]:
        return "总体结论: 不合规 —— 存在强制条款不满足，须整改后复审。"
    if summary["manual"]:
        return ("总体结论: 待人工确认 —— 自动核查未发现违规，"
                "但有 %d 项需人工复核。" % summary["manual"])
    return "总体结论: 通过（基于当前规则库）。"


def _findings_docx(document, findings) -> None:
    if not findings:
        document.add_paragraph("（无）")
        return
    table = document.add_table(rows=1, cols=6)
    table.style = "Light Grid Accent 1"
    _fill_row(table.rows[0], ("规则", "级别", "说明", "依据条款", "整改建议", "原文证据"))
    for f in sorted(findings, key=lambda x: (RESULT_ORDER.get(x.result, 9), x.rule_id)):
        _fill_row(table.add_row(), (
            "%s %s" % (f.rule_id, f.rule_name), f.level,
            "%s（%s）" % (f.detail or "-", RESULT_LABEL.get(f.result, f.result)),
            f.basis or "-", f.advice or "-", f.evidence_line or "-"))
