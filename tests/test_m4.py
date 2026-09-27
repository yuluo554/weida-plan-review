# -*- coding: utf-8 -*-
"""M4 测试：知识库检索问答、LLM 兜底（防幻觉三件套/降级）、基准脚本冒烟。

全部离线：LLM 客户端用假传输/假客户端，不打真实网络（.env 已配置也不走网络）。
"""
import io
import json
import sys
import tempfile
import unittest
from contextlib import redirect_stdout
from pathlib import Path
from unittest import mock

from planguard.ir.schema import Evidence, ParameterCard
from planguard.llm.client import (LLMConfig, LLMError, OpenAICompatClient,
                                  _loads_lenient, _strip_noise)
from planguard.parsers.base import ParseResult, Section, TextLine

REPO = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(REPO / "benchmarks"))

import eval_e2e  # noqa: E402
import eval_parse  # noqa: E402
from common import DEFAULT_GOLD_DIR, DEFAULT_SAMPLES_DIR  # noqa: E402

try:
    from planguard.parsers.docx_parser import HAS_DOCX
except ImportError:  # pragma: no cover
    HAS_DOCX = False


def has_docx():
    try:
        import docx  # noqa: F401
        return True
    except ImportError:
        return False


class TestKnowledgeStore(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from planguard.knowledge.store import load_blocks, load_thresholds
        cls.blocks = load_blocks(REPO / "data" / "knowledge" / "blocks")
        cls.thresholds = load_thresholds(REPO / "data" / "knowledge" / "thresholds")

    def test_data_scale(self):
        self.assertGreaterEqual(len(self.blocks), 30)
        self.assertGreaterEqual(len(self.thresholds), 20)
        for b in self.blocks:
            self.assertEqual(b.status, "待核对")

    def test_search_relevance(self):
        from planguard.knowledge.store import search
        results = search(self.blocks, "监测频率 开挖至坑底 加密")
        self.assertTrue(results)
        self.assertEqual(results[0][0].doc, "GB 50497-2019")
        results2 = search(self.blocks, "专家论证专家人数不得少于")
        self.assertIn("专家论证", results2[0][0].text)

    def test_match_threshold(self):
        from planguard.knowledge.store import match_threshold
        hit, _score = match_threshold(self.thresholds, "基坑开挖深度超过多少需要专家论证")
        self.assertIsNotNone(hit)
        self.assertEqual(hit.id, "T-DP-002")
        miss, _ = match_threshold(self.thresholds, "今天中午吃什么饭比较好")
        self.assertIsNone(miss)

    def test_ask_routes(self):
        from planguard.knowledge.store import ask
        kdir = REPO / "data" / "knowledge"
        r1 = ask("基坑开挖深度超过多少需要专家论证", kdir)
        self.assertEqual(r1["route"], "threshold")
        self.assertIn("专家论证", r1["answer"])
        r2 = ask("支撑轴力报警值怎么定", kdir)
        self.assertEqual(r2["route"], "blocks")
        self.assertIn("GB 50497-2019", r2["answer"])
        r3 = ask("量子力学的测不准原理是什么", kdir)
        self.assertEqual(r3["route"], "empty")

    def test_format_block_single_status(self):
        from planguard.knowledge.store import format_block
        block = next(b for b in self.blocks if b.doc == "GB 50497-2019")
        self.assertEqual(format_block(block).count("待核对"), 1)

    def test_cli_ask(self):
        from planguard.cli import main
        buf = io.StringIO()
        with redirect_stdout(buf):
            code = main(["ask", "基坑开挖深度超过多少需要专家论证"])
        self.assertEqual(code, 0)
        self.assertIn("专家论证", buf.getvalue())


def _client():
    return OpenAICompatClient(LLMConfig(base_url="http://fake", api_key="k", model="m",
                                        max_retries=2, timeout=5))


class TestLLMClient(unittest.TestCase):
    def test_json_repair(self):
        parse = lambda t: _loads_lenient(_strip_noise(t))  # noqa: E731
        self.assertEqual(parse('{"a": [1, 2}'), {"a": [1, 2]})
        self.assertEqual(parse('{"a": [1, 2,]}'), {"a": [1, 2]})
        self.assertEqual(parse('说明：{"a": 1} 以上。'), {"a": 1})
        fence = chr(96) * 3
        self.assertEqual(parse(fence + 'json\n{"items": []}\n' + fence), {"items": []})
        self.assertIsNone(parse('{"items": [完全不是json'))

    def test_retry_then_success(self):
        client = _client()
        responses = [TimeoutError("t"),
                     {"choices": [{"message": {"content": '{"items": [{"param": "x"}]'}}]}]

        def _fake_request(_payload):
            item = responses.pop(0)
            if isinstance(item, Exception):
                raise item
            return item

        with mock.patch.object(OpenAICompatClient, "_request", side_effect=_fake_request):
            out = client.complete_json("任意")
        self.assertEqual(out, {"items": [{"param": "x"}]})

    def test_no_retry_on_auth_error(self):
        import urllib.error
        client = _client()
        err = urllib.error.HTTPError("http://fake", 401, "Unauthorized", None, None)
        with mock.patch.object(OpenAICompatClient, "_request", side_effect=err) as mocked:
            with self.assertRaises(LLMError):
                client.complete("hello")
        self.assertEqual(mocked.call_count, 1)

    def test_exhausted_retries(self):
        client = _client()
        with mock.patch.object(OpenAICompatClient, "_request",
                               side_effect=TimeoutError("t")) as mocked:
            with self.assertRaises(LLMError):
                client.complete("hello")
        self.assertEqual(mocked.call_count, client.config.max_retries + 1)


class TestLLMFallback(unittest.TestCase):
    @staticmethod
    def _doc():
        return ParseResult(
            doc_name="fake.docx",
            sections=[Section(title="一、工程概况", level=1, path="一、工程概况")],
            lines=[
                TextLine(text="基坑开挖深度为6.2m，采用桩锚支护。", section_path="一、工程概况", page=3),
                TextLine(text="本段无参数的描述性文字。", section_path="一、工程概况", page=3),
                TextLine(text="计算书中支护结构顶部水平位移控制值取25mm。", section_path="一、工程概况", page=9),
            ],
        )

    @staticmethod
    def _missing():
        from planguard.extract.extractor import PARAM_META, TEXT_PARAMS
        return {"dp.excavation_depth": PARAM_META["dp.excavation_depth"],
                "dp.retaining.displacement": PARAM_META["dp.retaining.displacement"],
                "dp.safety_grade": TEXT_PARAMS["dp.safety_grade"]}

    def test_candidate_lines_compression(self):
        from planguard.llm.fallback import candidate_lines
        picked = candidate_lines(self._doc().lines, self._missing())
        texts = [ln.text for _no, ln in picked]
        self.assertNotIn("本段无参数的描述性文字。", texts)
        self.assertIn("基坑开挖深度为6.2m，采用桩锚支护。", texts)

    def test_validate_and_cards(self):
        from planguard.llm.fallback import fallback_extract

        class FakeClient:
            @staticmethod
            def complete_json(_prompt):
                return {"items": [
                    {"param": "dp.excavation_depth", "line_no": 0,
                     "excerpt": "基坑开挖深度为6.2m", "value": 6.2, "unit": "m"},
                    {"param": "dp.excavation_depth", "line_no": 0,
                     "excerpt": "基坑开挖深度为6.2m", "value": 7.7, "unit": "m"},
                    {"param": "dp.retaining.displacement", "line_no": 2,
                     "excerpt": "幻觉摘录不在原文里", "value": 25, "unit": "mm"},
                    {"param": "dp.retaining.displacement", "line_no": 2,
                     "excerpt": "位移控制值取25mm", "value": 900, "unit": "mm"},
                    {"param": "unknown.param", "line_no": 0, "excerpt": "基坑开挖深度为6.2m",
                     "value": 1, "unit": "m"},
                    {"param": "dp.safety_grade", "line_no": 0,
                     "excerpt": "基坑开挖深度为6.2m", "value": None, "text_value": "二级"},
                ]}

        cards, meta = fallback_extract(FakeClient(), self._doc(), self._missing())
        self.assertEqual(meta["cards_added"], 2)
        by_id = {c.param_id: c for c in cards}
        self.assertEqual(by_id["dp.excavation_depth"].value, 6.2)
        self.assertEqual(by_id["dp.excavation_depth"].source, "llm")
        self.assertEqual(by_id["dp.excavation_depth"].confidence, 0.6)
        self.assertEqual(by_id["dp.safety_grade"].text_value, "二级")
        reasons = " | ".join(d["reason"] for d in meta["dropped"])
        self.assertIn("子串", reasons)
        self.assertIn("范围", reasons)
        self.assertIn("缺失清单", reasons)


class TestOrchestratorLLM(unittest.TestCase):
    """编排器降级约定：未配置/失败都不影响纯规则结论。"""

    DOC = REPO / "data" / "samples" / "gen_0001.docx"

    @classmethod
    def setUpClass(cls):
        if not has_docx():
            raise unittest.SkipTest("python-docx 未安装")
        from planguard.orchestrator import run
        cls.pure, _ = run(cls.DOC)

    def test_unconfigured_degrades_cleanly(self):
        from planguard.orchestrator import run
        with mock.patch("planguard.llm.client.is_configured", return_value=False):
            result, _trace = run(self.DOC, use_llm=True)
        self.assertIn("纯规则", result.meta["llm"].get("reason", ""))
        self.assertEqual([f.result for f in result.findings],
                         [f.result for f in self.pure.findings])

    def test_llm_failure_degrades_cleanly(self):
        from planguard.extract.extractor import PARAM_META
        from planguard.orchestrator import run
        fake_missing = {"dp.monitoring.frequency": dict(PARAM_META["dp.monitoring.frequency"])}
        cfg = mock.patch("planguard.llm.client.is_configured", return_value=True)
        miss = mock.patch("planguard.orchestrator._missing_params",
                          return_value=fake_missing)
        cli = mock.patch("planguard.llm.client.OpenAICompatClient",
                         side_effect=LLMError("网络不通"))
        with cfg, miss, cli:
            result, _trace = run(self.DOC, use_llm=True)
        self.assertIn("LLMError", result.meta["llm"].get("degraded", ""))
        self.assertEqual(result.summary(), self.pure.summary())

    @unittest.skipUnless(has_docx(), "python-docx 未安装")
    def test_llm_success_adds_card(self):
        """缺监测方案的文档：LLM 兜底补出位移参数卡（假客户端，离线）。"""
        import docx
        from planguard.orchestrator import run

        document = docx.Document()
        document.add_heading("深基坑方案（缺监测参数，测 LLM 兜底）", level=0)
        document.add_heading("一、工程概况", level=1)
        document.add_paragraph("基坑开挖深度为6.0m。")
        document.add_paragraph("本工程基坑属超过一定规模的危大工程，已按规定组织专家论证。")
        document.add_heading("二、计算书", level=1)
        document.add_paragraph("水平位移按25毫米控制，计算书及图纸见附件。")
        with tempfile.TemporaryDirectory() as tmp:
            path = Path(tmp) / "no_monitor.docx"
            document.save(str(path))

            class FakeClient:
                @staticmethod
                def complete_json(_prompt):
                    return {"items": [
                        {"param": "dp.retaining.displacement", "line_no": 5,
                         "excerpt": "水平位移按25毫米控制", "value": 25, "unit": "mm"},
                    ]}

            with mock.patch("planguard.llm.client.is_configured", return_value=True), \
                 mock.patch("planguard.llm.client.OpenAICompatClient",
                            return_value=FakeClient()):
                result, trace = run(path, use_llm=True)
        llm_meta = result.meta["llm"]
        self.assertIn("dp.retaining.displacement", llm_meta.get("missing_params", []))
        self.assertEqual(llm_meta.get("cards_added"), 1)
        llm_cards = [c for c in result.cards if c.source == "llm"]
        self.assertEqual(len(llm_cards), 1)
        self.assertEqual(llm_cards[0].value, 25.0)
        self.assertEqual(llm_cards[0].evidence.line_text, "水平位移按25毫米控制，计算书及图纸见附件。")
        self.assertIn("llm_fallback", [st.name for st in trace.stages])


class TestBenchmarkSmoke(unittest.TestCase):
    """基准冒烟（plan/03 §8）：全量 30 份语料，指标达标 + 诊断落盘。"""

    def test_parse_benchmark(self):
        with tempfile.TemporaryDirectory() as tmp:
            res = eval_parse.evaluate(DEFAULT_SAMPLES_DIR, DEFAULT_GOLD_DIR,
                                      diag_path=Path(tmp) / "diag.jsonl")
        self.assertGreaterEqual(res["f1"], 0.95)
        self.assertEqual(res["docs_total"], 30)

    def test_e2e_benchmark(self):
        with tempfile.TemporaryDirectory() as tmp:
            res = eval_e2e.evaluate(DEFAULT_SAMPLES_DIR, DEFAULT_GOLD_DIR,
                                    diag_path=Path(tmp) / "diag.jsonl")
        self.assertEqual(res["hits"], res["injected_total"])
        self.assertEqual(res["forced_fp"], 0)
        self.assertFalse(res["doc_errors"])


if __name__ == "__main__":
    unittest.main()
